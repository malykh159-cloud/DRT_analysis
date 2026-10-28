from dataclasses import replace
import numpy as np
import pytest
from drt_analyzer.analysis.peaks import manual_peaks
from drt_analyzer.analysis.cpe import estimate_cpe
from drt_analyzer.analysis.pipeline import analyze
from drt_analyzer.drt.base import AnalysisSettings
from drt_analyzer.export.tables import summary_rows, export_xlsx
from openpyxl import load_workbook


def test_manual_regions_integrate_and_exclude_without_changing_total(rc):
    v = analyze(rc, AnalysisSettings(lambda_method="Manual", validation=False))
    split = float(v.tau[len(v.tau)//2])
    regions = [(np.sqrt(v.tau[0]*split), v.tau[0], split, False),
               (np.sqrt(split*v.tau[-1]), split, v.tau[-1], True)]
    peaks = manual_peaks(v.tau, v.gamma, regions)
    assert sum(p.rp_ohm for p in peaks) == pytest.approx(v.rp_total)
    edited = replace(v, peaks=peaks)
    assert edited.rp_total == v.rp_total
    np.testing.assert_array_equal(edited.gamma, v.gamma)
    rows = summary_rows([(rc, edited)])
    assert rows[1][4:] == [peaks[1].rp_ohm]
    with pytest.raises(ValueError, match="overlap"):
        manual_peaks(v.tau, v.gamma, [regions[0], (split, v.tau[0], split, True)])


def test_cpe_projection_known_rc():
    # A concentrated RC distribution should give n=1, Q=C=tau/R.
    from drt_analyzer.analysis.peaks import Peak
    tau = np.geomspace(.00999, .01001, 501)
    gamma = np.ones(501)
    area = np.trapezoid(gamma, np.log(tau))
    peak = Peak(1, .01, 1/(2*np.pi*.01), 1., area, 1., tau[0], tau[-1])
    q, n, note = estimate_cpe(tau, gamma, peak, np.geomspace(.01, 1e5, 81))
    assert n == pytest.approx(1., abs=1e-5)
    assert q == pytest.approx(.01/area, rel=1e-4)
    assert "projection" in note
    assert estimate_cpe(tau, gamma, replace(peak, edge=True), [1.])[0] is None


def test_export_independent_axes_and_duplicate_names(tmp_path, rc):
    v = analyze(rc, AnalysisSettings(lambda_method="Manual", validation=False,
                                    frequency_min_hz=1., frequency_max_hz=1000.))
    path = tmp_path/"compact.xlsx"
    export_xlsx(path, [(rc, v), (rc, v)])
    wb = load_workbook(path, data_only=True)
    assert len(wb.sheetnames) == 3
    s = wb.worksheets[1]
    assert sum(s.cell(i, 1).value is not None for i in range(2, s.max_row+1)) == len(rc.frequency)
    assert sum(s.cell(i, 6).value is not None for i in range(2, s.max_row+1)) == len(v.frequency)
    assert AnalysisSettings().lambda_min_exp == -6
    wb.close()


def test_cpe_projection_known_distributed_zarc():
    from drt_analyzer.analysis.peaks import Peak
    r, t0, n = 50., .01, .8
    tau = np.geomspace(1e-12, 1e8, 4001)
    x = np.log(tau/t0)
    gamma = r/(2*np.pi)*np.sin(np.pi*n)/(np.cosh(n*x)+np.cos(np.pi*n))
    area = float(np.trapezoid(gamma, np.log(tau)))
    p = Peak(1, t0, 1/(2*np.pi*t0), float(gamma.max()), area, 1., tau[0], tau[-1])
    q, fitted_n, _ = estimate_cpe(tau, gamma, p, np.geomspace(.01, 1e5, 81))
    assert fitted_n == pytest.approx(n, abs=1e-4)
    assert q == pytest.approx(t0**n/r, rel=1e-4)
