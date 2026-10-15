from dataclasses import replace
import json
from pathlib import Path
import numpy as np
import pytest
from openpyxl import load_workbook
from drt_analyzer.analysis.pipeline import analyze
from drt_analyzer.drt.base import AnalysisSettings
from drt_analyzer.state.history import History
from drt_analyzer.export.tables import export_xlsx, export_txt_bundle, summary_rows, safe_cell
from drt_analyzer.export.settings import export_settings


def test_history(rc):
    a = analyze(rc, AnalysisSettings(lambda_method="Manual", lambda_value=1e-6, validation=False))
    b = analyze(rc, replace(a.settings, lambda_value=.001))
    c = analyze(rc, replace(a.settings, lambda_value=.1))
    h = History()
    h.push(a)
    h.push(b)
    assert h.undo() is a
    assert h.redo() is b
    h.undo()
    h.push(c)
    assert h.current is c
    assert h.variants == [a, c]


def test_remove_history_preserves_current_and_neighbors():
    a, b, c, d = [object() for _ in range(4)]
    h = History([a, b, c, d], 2)
    assert h.remove([0, 3]) is c
    assert h.position == 1 and h.variants == [b, c]
    assert h.remove([1]) is b
    assert h.remove([0]) is None
    assert h.position == -1 and h.undo() is None and h.redo() is None
    h.push(d)
    assert h.current is d


def test_remove_history_invalid_selection():
    h = History([object()], 0)
    with pytest.raises(IndexError):
        h.remove([1])


def test_exports(tmp_path, rc):
    v = analyze(rc, AnalysisSettings(lambda_method="Manual", lambda_value=1e-6))
    export_xlsx(tmp_path/"result.xlsx", [(rc, v)])
    export_txt_bundle(tmp_path/"txt", [(rc, v)])
    export_settings(tmp_path/"settings.json", rc, v)
    m = json.loads((tmp_path/"settings.json").read_text(encoding="utf-8"))
    assert m["source_sha256"] == rc.source_hash
    assert m["actual_lambda"] == v.lambda_value
    assert m["settings"]["derivative_order"] == 2
    assert m["preprocessing"]["points_used"] == len(rc.frequency)
    assert m["import"]["im_multiplier"] == 1
    wb = load_workbook(tmp_path/"result.xlsx", read_only=True, data_only=True)
    assert wb["Summary"].cell(2, 3).value == v.rp_total
    assert wb["1 DRT"].max_row == len(v.tau)+1
    assert wb["1 Raw source"].cell(2, 4).value == pytest.approx(rc.imaginary_original[0], rel=1e-14)
    assert wb["1 Impedance"].max_row == len(rc.frequency)+1
    wb.close()
    assert len(list((tmp_path/"txt").glob("*.txt"))) == 7
    assert safe_cell("=1+1") == "'=1+1"


def test_deterministic_auto(rc):
    settings = AnalysisSettings(lambda_count=9)
    a, b = analyze(rc, settings), analyze(rc, settings)
    assert a.lambda_value == b.lambda_value
    np.testing.assert_array_equal(a.gamma, b.gamma)
