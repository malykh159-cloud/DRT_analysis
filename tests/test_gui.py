from dataclasses import replace
import os
import time
from pathlib import Path
import numpy as np
import pytest
from PySide6 import QtCore, QtGui, QtTest, QtWidgets as QtW
from drt_analyzer.app.main_window import MainWindow, Entry
from drt_analyzer.app.main_window import variant_status
from drt_analyzer.analysis.pipeline import analyze
from drt_analyzer.drt.base import AnalysisSettings
from drt_analyzer.ui.lambda_scan import LambdaScanDialog
from drt_analyzer.analysis.pipeline import lambda_scan
from drt_analyzer.export.plots import export_plot
from drt_analyzer.ui.parameter_help import PARAMETER_HELP


@pytest.fixture(scope="module")
def app():
    instance = QtW.QApplication.instance() or QtW.QApplication([])
    if not QtGui.QFontDatabase.families():
        QtGui.QFontDatabase.addApplicationFont("C:/Windows/Fonts/segoeui.ttf")
    instance.setFont(QtGui.QFont("Segoe UI", 9))
    return instance


def wait_until(app, predicate, timeout=15):
    end = time.monotonic()+timeout
    while not predicate():
        app.processEvents()
        if time.monotonic() > end:
            raise AssertionError("GUI worker timed out")
        time.sleep(.005)
    app.processEvents()


def test_manual_peak_editor_exclude_add_and_history(app, rc, monkeypatch):
    from drt_analyzer.ui.peak_editor import PeakEditor
    v = analyze(rc, AnalysisSettings(lambda_method="Manual", validation=False))
    dialog = PeakEditor(rc, v)
    dialog.table.item(0, 0).setCheckState(QtCore.Qt.CheckState.Unchecked)
    assert not dialog.result_peaks[0].included
    assert "Selected Rp = 0" in dialog.total.text()
    marker = np.sqrt(v.peaks[0].tau_s*v.tau[-1])
    monkeypatch.setattr(QtW.QInputDialog, "getText", lambda *args: (str(1/(2*np.pi*marker)), True))
    dialog.add_marker()
    assert len(dialog.result_peaks) == 2
    assert sum(p.included for p in dialog.result_peaks) == 1
    assert sum(p.rp_ohm for p in dialog.result_peaks) == pytest.approx(v.rp_total)
    dialog.accept()
    assert dialog.result() == QtW.QDialog.DialogCode.Accepted
    window = MainWindow()
    window.entries.append(Entry(rc))
    window.refresh_table()
    window.table.setCurrentCell(0, 0)
    window.result_ready(0, v)
    monkeypatch.setattr(PeakEditor, "exec", lambda self: (self.table.item(0, 0).setCheckState(
        QtCore.Qt.CheckState.Unchecked), self.accept(), QtW.QDialog.DialogCode.Accepted)[-1])
    window.edit_peaks()
    assert not window.current.history.current.peaks[0].included
    assert window.current.history.current.rp_total == v.rp_total
    window.undo()
    assert window.current.history.current is v
    window.redo()
    assert not window.current.history.current.peaks[0].included
    scan = LambdaScanDialog(rc, (v,), window)
    scan.edit_peaks()
    assert not scan.chosen.peaks[0].included
    assert v.peaks[0].included
    scan.close()
    dialog.close()
    window.close()


def test_window_history_plots_and_export(app, rc, tmp_path):
    window = MainWindow()
    window.show()
    window.entries.append(Entry(rc))
    window.refresh_table()
    window.table.setCurrentCell(0, 0)
    settings = AnalysisSettings(lambda_method="Manual", lambda_value=1e-6, validation=False)
    a, b = analyze(rc, settings), analyze(rc, replace(settings, lambda_value=.001))
    window.result_ready(0, a)
    window.result_ready(0, b)
    window.undo()
    assert window.current.history.current is a
    assert window.current.status == variant_status(a)
    window.redo()
    assert window.current.history.current is b
    window.settings.lambda_method.setCurrentText("Manual")
    window.settings.lambda_value.setText("0.02")
    assert window.settings.value().lambda_value == .02
    window.plots.axis.setCurrentIndex(1)
    window.plots.compare([(rc, b)])
    app.processEvents()
    for name, plot in window.plots.all_plots().items():
        for suffix in ("png",):
            path = tmp_path/f"{name}.{suffix}"
            export_plot(plot, path)
            assert path.stat().st_size > 500
    window.close()


def test_background_result_and_cancellation(app, rc):
    window = MainWindow()
    window.entries = [Entry(rc), Entry(rc)]
    window.refresh_table()
    window.table.setCurrentCell(0, 0)
    settings = AnalysisSettings(lambda_method="Manual", lambda_value=1e-6, validation=False)
    window.start_jobs([(0, rc, settings), (1, rc, settings)])
    wait_until(app, lambda: window.thread is None)
    assert all(e.history.current for e in window.entries)
    assert window.progress.value() == 2
    old = window.entries[0].history.current
    slow = replace(settings, lambda_method="k-fold", lambda_count=81)
    window.start_jobs([(0, rc, slow), (1, rc, slow)])
    window.cancel()
    wait_until(app, lambda: window.thread is None)
    assert window.entries[0].history.current is old
    assert window.entries[1].status == "Cancelled"
    window.close()


def test_scan_slider_uses_cached_objects(app, rc):
    settings = AnalysisSettings(lambda_method="Manual", lambda_count=7, validation=False)
    variants = lambda_scan(rc, settings)
    dialog = LambdaScanDialog(rc, variants)
    dialog.slider.setValue(4)
    assert dialog.chosen is variants[4]
    assert dialog.plot.variant is variants[4]
    dialog.close()


def test_manual_range_controls_and_auto_preserve_per_file(app, rc, monkeypatch):
    window = MainWindow()
    window.entries = [Entry(rc), Entry(rc, draft=AnalysisSettings(frequency_min_hz=1., frequency_max_hz=1e3))]
    window.refresh_table()
    window.table.setCurrentCell(0, 0)
    assert window.calculate_action.text() == "Run analysis"
    window.settings.frequency_range.setChecked(True)
    window.settings.frequency_min.setText("0,1")
    window.settings.frequency_max.setText("1e4")
    jobs = []
    monkeypatch.setattr(window, "start_jobs", lambda pending: jobs.extend(pending))
    window.auto_all()
    assert jobs[0][2].frequency_min_hz == .1 and jobs[0][2].frequency_max_hz == 1e4
    assert jobs[1][2].frequency_min_hz == 1 and jobs[1][2].frequency_max_hz == 1e3
    assert all(settings.mode == "Auto" for _, _, settings in jobs)
    window.table.setCurrentCell(1, 0)
    assert window.settings.value().frequency_min_hz == 1
    window.table.setCurrentCell(0, 0)
    assert window.settings.value().frequency_min_hz == .1
    window.settings.frequency_range.setChecked(False)
    assert window.settings.value().frequency_min_hz is None
    window.close()


def test_del_removes_spectra_without_deleting_sources(app, tmp_path):
    paths = []
    for i in range(3):
        path = tmp_path/f"source{i}.csv"
        path.write_text("f;Re;Im\n"+"\n".join(f"{j+1};{i+10};-{j+1}" for j in range(10)))
        paths.append(path)
    original = [p.read_bytes() for p in paths]
    window = MainWindow()
    window.show()
    window.add_files(paths)
    window.table.setCurrentCell(1, 0)
    survivor = window.current
    window.table.clearSelection()
    flags = QtCore.QItemSelectionModel.SelectionFlag.Select | QtCore.QItemSelectionModel.SelectionFlag.Rows
    for row in (0, 2):
        window.table.selectionModel().select(window.table.model().index(row, 0), flags)
    window.activateWindow()
    window.table.setFocus()
    app.processEvents()
    QtTest.QTest.keyClick(window.table, QtCore.Qt.Key.Key_Delete)
    assert len(window.entries) == 1 and window.current is survivor
    assert [p.read_bytes() for p in paths] == original
    window.settings.lambda_value.setFocus()
    window.settings.lambda_value.selectAll()
    QtTest.QTest.keyClick(window.settings.lambda_value, QtCore.Qt.Key.Key_Delete)
    assert len(window.entries) == 1  # Del in an editor must not delete a spectrum
    window.table.setFocus()
    window.table.selectAll()
    QtTest.QTest.keyClick(window.table, QtCore.Qt.Key.Key_Delete)
    assert not window.entries and window.current is None
    assert window.history_list.count() == 0 and window.peaks.rowCount() == 0
    assert not window.plots._pairs
    window.close()


def test_del_removes_multiple_history_entries_and_refreshes_compare(app, rc):
    window = MainWindow()
    window.show()
    window.entries = [Entry(rc)]
    window.refresh_table()
    window.table.setCurrentCell(0, 0)
    settings = AnalysisSettings(lambda_method="Manual", validation=False)
    variants = [analyze(rc, replace(settings, lambda_value=lam)) for lam in (1e-6, 1e-4, 1e-2)]
    for v in variants:
        window.result_ready(0, v)
    window.plots.compare([(rc, variants[-1])])
    window.history_list.clearSelection()
    window.history_list.item(0).setSelected(True)
    window.history_list.item(2).setSelected(True)
    window.activateWindow()
    window.history_list.setFocus()
    app.processEvents()
    QtTest.QTest.keyClick(window.history_list, QtCore.Qt.Key.Key_Delete)
    assert window.current.history.variants == [variants[1]]
    assert window.plots.variant is variants[1]
    assert window.plots._pairs[0][1] is variants[1]
    window.history_list.selectAll()
    QtTest.QTest.keyClick(window.history_list, QtCore.Qt.Key.Key_Delete)
    assert window.current.history.current is None
    assert window.current.status == "Imported"
    assert window.plots.variant is None and not window.plots._pairs
    assert window.peaks.rowCount() == 0 and not window.undo_action.isEnabled()
    assert len(window.entries) == 1
    window.close()


def test_parameter_help_buttons_are_complete_and_modeless(app):
    window = MainWindow()
    keys = {button.objectName()[5:] for button in window.settings.findChildren(QtW.QToolButton)
            if button.objectName().startswith("help_")}
    assert keys == set(PARAMETER_HELP)-{"drt_axis"}
    for key in sorted(keys):
        button = window.settings.findChild(QtW.QToolButton, "help_"+key)
        assert button.isEnabled()
        button.click()
        assert button.help_dialog.isVisible()
        assert not button.help_dialog.isModal()
        text = button.help_dialog.text.toPlainText()
        assert all(title in text for title in ("Что это", "На что влияет", "Практический выбор"))
        button.help_dialog.close()
    window.close()


def test_frequency_readouts_and_hidden_comparison_curves(app, rc):
    window = MainWindow()
    window.show()
    window.entries = [Entry(rc)]
    window.refresh_table()
    window.table.setCurrentCell(0, 0)
    v = analyze(rc, AnalysisSettings(lambda_method="Manual", validation=True))
    window.result_ready(0, v)
    app.processEvents()
    plots = window.plots
    # Use a well-separated point on the arc for a physical mouse event: dense
    # HF samples can be closer than one integer screen pixel after projection.
    i = 43
    scene_pos = plots.nyquist.plotItem.vb.mapViewToScene(QtCore.QPointF(rc.real[i], -rc.impedance[i].imag))
    hover = plots.hovers[plots.nyquist]
    window.activateWindow()
    QtTest.QTest.mouseMove(plots.nyquist.viewport(), plots.nyquist.mapFromScene(scene_pos))
    app.processEvents()
    assert hover.cursor_frequency == rc.frequency[i]
    assert "Ближайшая точка" in plots.cursor_readout.text()
    plots.tabs.setCurrentIndex(2)
    app.processEvents()
    hover = plots.hovers[plots.drt]
    scene_pos = plots.drt.plotItem.vb.mapViewToScene(QtCore.QPointF(np.log10(.01), max(v.gamma)/2))
    hover.mouse_moved(scene_pos)
    assert hover.cursor_frequency == pytest.approx(1/(2*np.pi*.01))
    assert "τ =" in plots.cursor_readout.text()
    plots.axis.setCurrentIndex(1)
    app.processEvents()
    scene_pos = plots.drt.plotItem.vb.mapViewToScene(QtCore.QPointF(np.log10(100.), max(v.gamma)/2))
    hover.mouse_moved(scene_pos)
    assert hover.cursor_frequency == pytest.approx(100.)
    plots.compare([(rc, v)])
    app.processEvents()
    compare = plots.hovers[plots.comparison]
    scene_pos = plots.comparison.plotItem.vb.mapViewToScene(QtCore.QPointF(np.log10(100.), max(v.gamma)/2))
    compare.mouse_moved(scene_pos)
    assert compare.cursor_frequency == pytest.approx(100.)
    compare.curves[0].item.hide()
    compare.mouse_moved(scene_pos)
    assert compare.cursor_frequency is None
    window.close()


def test_editing_auto_values_switches_only_corresponding_mode(app):
    from drt_analyzer.ui.settings_panel import SettingsPanel
    panel = SettingsPanel()
    assert panel.shape_mode.currentText() == "Manual" and panel.shape.value() == 6
    panel.load(AnalysisSettings(mode="Auto", lambda_method="rGCV", shape_mode="Auto"))
    assert panel.lambda_value.isEnabled() and panel.shape.isEnabled()
    assert panel.frequency_min.isEnabled() and panel.frequency_max.isEnabled()
    panel.lambda_value.selectAll()
    QtTest.QTest.keyClicks(panel.lambda_value, "0.02")
    assert panel.lambda_method.currentText() == "Manual"
    assert panel.shape_mode.currentText() == "Auto"
    panel.shape.setValue(7)
    assert panel.shape_mode.currentText() == "Manual"
    panel.frequency_min.selectAll()
    QtTest.QTest.keyClicks(panel.frequency_min, "1")
    assert panel.frequency_range.isChecked()
    panel.load(AnalysisSettings(shape_mode="Auto"))
    assert panel.shape_mode.currentText() == "Auto"  # loading history is not a user edit
    panel.close()


def test_bulk_export_all_formats_and_graph_choices(app, rc, tmp_path):
    from drt_analyzer.export.batch import export_batch
    from drt_analyzer.ui.export_dialog import ExportDialog, PLOT_TYPES
    from openpyxl import load_workbook
    dialog = ExportDialog(2, 1)
    assert dialog.scope.currentIndex() == 0
    assert dialog.plot_types() == ["Nyquist", "DRT"]
    a = analyze(rc, AnalysisSettings(lambda_method="Manual", lambda_value=1e-6))
    b = analyze(rc, replace(a.settings, lambda_value=.01))
    destination = tmp_path/"batch"
    export_batch(destination, [(rc, a), (rc, b)], txt=True, plot_types=PLOT_TYPES)
    pngs = list((destination/"plots").glob("*.png"))
    assert len(pngs) == 12 and all(p.stat().st_size > 2000 for p in pngs)
    assert (destination/"plots"/"001_synthetic_DRT.png").read_bytes() != (destination/"plots"/"002_synthetic_DRT.png").read_bytes()
    assert {p.suffix for p in destination.rglob("*") if p.is_file()} == {".xlsx", ".txt", ".png"}
    wb = load_workbook(destination/"results.xlsx", read_only=True, data_only=True)
    assert wb["Summary"].max_row == 3
    assert [c.value for c in wb["Summary"][1]][:4] == ["File", "Rs (Ohm)", "Rp total (Ohm)", "Lambda"]
    wb.close()
    dialog.close()
