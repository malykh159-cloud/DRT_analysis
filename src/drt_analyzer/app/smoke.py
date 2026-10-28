"""Executable-level smoke verification: synthetic numerical truth + real Qt export."""
from dataclasses import replace
import json
import sys
from pathlib import Path
import numpy as np
from PySide6 import QtCore, QtWidgets as QtW
from ..data.models import RawSpectrum, ImportInfo
from ..drt.base import AnalysisSettings, BayesianSettings
from ..analysis.pipeline import analyze, lambda_scan
from ..export.tables import export_xlsx, export_txt_bundle
from ..export.settings import export_settings
from ..export.plots import export_plot
from ..export.batch import export_batch
from ..data.importers import load_spectrum
from ..ui.export_dialog import ExportDialog
from .main_window import Entry
from .. import __version__


def run_smoke(window, output):
    output = Path(output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    f = np.geomspace(1e5, .01, 81)
    z = 5+30/(1+2j*np.pi*f*1e-4)+70/(1+2j*np.pi*f*.1)+2j*np.pi*f*1e-5
    raw = RawSpectrum("Synthetic two RC validation.csv", "synthetic-not-a-file-hash", f, z.real, z.imag,
        ImportInfo("synthetic", ",", ("Frequency", "Re", "Im"), (0, 1, 2), "Im(Z)", 1, "synthetic truth"))
    settings = AnalysisSettings(lambda_method="Manual", lambda_value=1e-6)
    v = analyze(raw, settings)
    assert len(v.peaks) == 2
    assert abs(v.rp_total-100) < 2
    assert v.quality["relative_rms"] < .01
    window.entries.append(Entry(raw))
    window.refresh_table()
    window.table.setCurrentCell(0, 0)
    window.table.selectRow(0)
    window.result_ready(0, v)
    second = analyze(raw, replace(settings, lambda_value=1e-4))
    window.result_ready(0, second)
    window.undo()
    assert window.current.history.current is v
    window.redo()
    assert window.current.history.current is second
    window.undo()
    window.plots.tabs.setCurrentIndex(2)
    window.plots.compare([(raw, v)], switch=False)
    QtW.QApplication.processEvents()
    window.grab().save(str(output/"workspace.png"))
    for name, item in window.plots.all_plots().items():
        for suffix in ("png",):
            export_plot(item, output/f"{name}.{suffix}")
    export_xlsx(output/"results.xlsx", [(raw, v)])
    # Exercise the compact workbook contract and new editor in the frozen runtime.
    from openpyxl import load_workbook
    from ..ui.peak_editor import PeakEditor
    book = load_workbook(output/"results.xlsx", read_only=True)
    assert len(book.sheetnames) == 2 and book.sheetnames[0] == "Summary"
    assert book["Summary"].cell(1, 6).value == "Peak 1 f_max (Hz)"
    book.close()
    editor = PeakEditor(raw, v, window)
    editor.show()
    editor.table.item(0, 0).setCheckState(QtCore.Qt.CheckState.Unchecked)
    assert not editor.result_peaks[0].included
    assert editor.result_peaks[1].included
    QtW.QApplication.processEvents()
    editor.grab().save(str(output/"manual_peaks.png"))
    editor.accept()
    editor.close()
    export_txt_bundle(output/"txt", [(raw, v)])
    export_settings(output/"settings.json", raw, v)
    # Verify 0.3.0 paths in the frozen runtime, including headerless TXT import.
    source = output/"headerless.txt"
    np.savetxt(source, np.column_stack((f, z.real, z.imag)), delimiter="\t")
    imported = load_spectrum(source)
    np.testing.assert_array_equal(imported.frequency, f)
    assert imported.info.header_row == 0
    assert window.settings.shape_mode.currentText() == "Manual" and window.settings.shape.value() == 6
    window.settings.lambda_method.setCurrentText("rGCV")
    window.settings.lambda_value.textEdited.emit("0.001")
    assert window.settings.lambda_method.currentText() == "Manual"
    dialog = ExportDialog(2, 1, window)
    assert dialog.plot_types() == ["Nyquist", "DRT"]
    export_batch(output/"batch", [(raw, v), (raw, second)], txt=True, plot_types=dialog.plot_types())
    assert len(list((output/"batch"/"plots").glob("*.png"))) == 4
    dialog.close()
    # Exercise the six 0.2.0 user-facing additions in the frozen executable too.
    ranged = analyze(raw, replace(settings, frequency_min_hz=float(f[-8]), frequency_max_hz=float(f[5])))
    window.result_ready(0, ranged)
    assert len(ranged.frequency) < len(raw.frequency)
    assert window.calculate_action.text() == "Run analysis"
    window.plots.tabs.setCurrentIndex(2)
    QtW.QApplication.processEvents()
    hover = window.plots.hovers[window.plots.drt]
    scene = window.plots.drt.plotItem.vb.mapViewToScene(QtCore.QPointF(np.log10(.01), max(ranged.gamma)/2))
    hover.mouse_moved(scene)
    assert np.isclose(hover.cursor_frequency, 1/(2*np.pi*.01))
    window.grab().save(str(output/"workspace_range.png"))
    window.plots.tabs.setCurrentIndex(0)
    QtW.QApplication.processEvents()
    scene = window.plots.nyquist.plotItem.vb.mapViewToScene(QtCore.QPointF(raw.real[0], -raw.impedance[0].imag))
    window.plots.hovers[window.plots.nyquist].mouse_moved(scene)
    assert window.plots.hovers[window.plots.nyquist].cursor_frequency == raw.frequency[0]
    window.grab().save(str(output/"nyquist_hover.png"))
    button = window.settings.findChild(QtW.QToolButton, "help_lambda_value")
    button.click()
    QtW.QApplication.processEvents()
    button.help_dialog.grab().save(str(output/"help_lambda.png"))
    button.help_dialog.close()
    export_settings(output/"range_settings.json", raw, ranged)
    export_xlsx(output/"range_results.xlsx", [(raw, ranged)])
    from ..data.column_detection import imaginary_sign
    assert imaginary_sign("Im", np.array([1, 2, -1]))[0] == -1
    assert imaginary_sign("-Im", np.array([-1, -2, 1]))[0] == 1
    window.remove_selected_history()
    assert window.current.history.current is v
    window.remove_selected_spectra()
    assert window.current is None and not window.entries
    # Exercise the real Bayesian solver, credible plot and history in the bundle.
    bf = np.geomspace(1e3, .1, 21)
    bz = 5+80/(1+2j*np.pi*bf*.01)
    braw = RawSpectrum("Bayesian smoke.txt", "synthetic", bf, bz.real, bz.imag, raw.info)
    bs = AnalysisSettings(method="Bayesian DRT", lambda_method="Manual", lambda_value=.001,
        validation=False, bayesian=BayesianSettings(warmup=30, samples=80, chains=2))
    bv = analyze(braw, bs)
    assert abs(bv.rp_total/80-1) < .08 and len(bv.peaks) == 1
    assert bv.uncertainty.coefficient_draws.shape[:2] == (2, 80)
    window.entries.append(Entry(braw))
    window.refresh_table()
    window.table.setCurrentCell(0, 0)
    window.result_ready(0, bv)
    window.plots.tabs.setCurrentIndex(2)
    QtW.QApplication.processEvents()
    window.grab().save(str(output/"bayesian_workspace.png"))
    export_plot(window.plots.drt.plotItem, output/"bayesian_drt.png")
    export_xlsx(output/"bayesian_results.xlsx", [(braw, bv)])
    scan_events = []
    scan_settings = replace(bs, lambda_min_exp=-3, lambda_max_exp=-2, lambda_count=7,
                            bayesian=BayesianSettings(warmup=0, samples=20, chains=1))
    scanned = lambda_scan(braw, scan_settings,
        sampling_progress=lambda done, total, message: scan_events.append((done, total, message)))
    assert len(scanned) == 7 and all(item.uncertainty is not None for item in scanned)
    assert any(done > 0 and "λ 1/7" in message and "Chain 1/1" in message
               for done, total, message in scan_events)
    assert scan_events[-1][0] == scan_events[-1][1]
    result = {"passed": True, "software_version": __version__, "frozen": bool(getattr(sys, "frozen", False)),
              "python_executable": sys.executable, "rp_total": v.rp_total, "peaks": len(v.peaks),
              "relative_rms": v.quality["relative_rms"], "kk_status": v.kk.status,
              "lambda": v.lambda_value, "shape": v.shape_value,
              "undo_redo": "passed", "plots": 7, "frequency_range_points_used": len(ranged.frequency),
              "sign_majority": "passed", "delete_history": "passed", "delete_spectra": "passed",
              "parameter_help": "passed", "frequency_hover": "passed"}
    result.update(headerless_import="passed", bulk_export="passed", manual_settings="passed")
    result.update(bayesian_calculation="passed", bayesian_rp=bv.rp_total,
                  bayesian_diagnostics=bv.uncertainty.diagnostics)
    result["bayesian_scan_live_progress"] = "passed"
    (output/"smoke_result.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
