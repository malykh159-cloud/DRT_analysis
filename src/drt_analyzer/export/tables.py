import csv
import re
from pathlib import Path
from .settings import metadata


def safe_cell(value):
    # Prevent filenames and free-text headers becoming executable Excel formulas.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'"+value
    return value


def write_txt(path, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream, delimiter="\t")
        writer.writerows([[safe_cell(v) for v in row] for row in rows])


def summary_rows(pairs, cpe=False):
    max_peaks = max((sum(p.included for p in v.peaks) for _, v in pairs), default=0)
    header = ["File", "Rs (Ohm)", "Rp total (Ohm)", "Lambda"]
    for i in range(1, max_peaks+1):
        header.append(f"Peak {i} Rp (Ohm)")
        if cpe:
            header.extend([f"Peak {i} f_max (Hz)", f"Peak {i} Q (S s^n, approx)", f"Peak {i} n (approx)"])
    rows = [header]
    for raw, v in pairs:
        row = [raw.name, v.rs_ohm, v.rp_total, v.lambda_value]
        for peak in v.peaks:
            if not peak.included:
                continue
            row.append(peak.rp_ohm)
            if cpe:
                from ..analysis.cpe import estimate_cpe
                q, n, _ = estimate_cpe(v.tau, v.gamma, peak, v.frequency)
                row.extend([peak.frequency_hz, q, n])
        rows.append(row+[""]*(len(header)-len(row)))
    return rows


def detail_tables(raw, v):
    drt = [["Tau (s)", "Frequency (Hz)", "Gamma (Ohm per d ln tau)"]]
    import numpy as np
    drt.extend(zip(v.tau, 1/(2*np.pi*v.tau), v.gamma))
    z = [["Frequency (Hz)", "Re experimental (Ohm)", "Im experimental (Ohm)",
          "Re DRT (Ohm)", "Im DRT (Ohm)", "Relative residual Re", "Relative residual Im",
          "Re KK (Ohm)", "Im KK (Ohm)", "KK residual Re", "KK residual Im"]]
    for i, f in enumerate(v.frequency):
        row = [f, v.impedance[i].real, v.impedance[i].imag, v.reconstructed[i].real,
               v.reconstructed[i].imag, v.residuals[i].real, v.residuals[i].imag]
        row += [v.kk.reconstructed[i].real, v.kk.reconstructed[i].imag,
                v.kk.residuals[i].real, v.kk.residuals[i].imag] if v.kk else [None]*4
        z.append(row)
    peaks = [["Index", "Tau (s)", "Frequency (Hz)", "Gamma peak (Ohm)", "Rp (Ohm)",
              "Fraction", "Left tau (s)", "Right tau (s)", "Unresolved edge", "Included"]]
    for p in v.peaks:
        peaks.append([p.index, p.tau_s, p.frequency_hz, p.gamma_ohm, p.rp_ohm,
                      p.fraction, p.left_tau_s, p.right_tau_s, p.edge, p.included])
    source = [["Source row (1-based)", "Frequency (Hz)", "Re original (Ohm)", "Imaginary original (Ohm)"]]
    source.extend((i+1, f, re, im) for i, (f, re, im) in enumerate(zip(raw.frequency, raw.real, raw.imaginary_original)))
    selection = [["Lambda", "Selection score", "Residual norm (L-curve)", "Penalty norm (L-curve)"]]
    sel = v.lambda_selection
    for i, lam in enumerate(sel.lambdas):
        selection.append([lam, float(sel.scores[i]) if i < len(sel.scores) and np.isfinite(sel.scores[i]) else None,
                          float(sel.residual_norms[i]) if i < len(sel.residual_norms) else None,
                          float(sel.penalty_norms[i]) if i < len(sel.penalty_norms) else None])
    return {"DRT": drt, "Impedance": z, "Peaks": peaks, "Raw source": source, "Lambda selection": selection}


def _format_sheet(ws):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244A64")
    for column in ws.columns:
        length = max(len(str(c.value or "")) for c in list(column)[:100])
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(60, max(15, length+2))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = "0.000000E+00"


def export_xlsx(path, pairs):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from ..analysis.cpe import estimate_cpe
    import numpy as np
    pairs = list(pairs)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for row in summary_rows(pairs, cpe=True):
        ws.append([safe_cell(v) for v in row])
    for n, (raw, v) in enumerate(pairs, 1):
        selected = [p for p in v.peaks if p.included]
        ws.cell(n+1, 3).comment = Comment(
            f"Full finite-window DRT integral, including excluded peaks. Selected Rp sum: "
            f"{sum(p.rp_ohm for p in selected):.12g} Ohm. Excluded original peak indices: "
            f"{[p.index for p in v.peaks if not p.included]}", "DRT Analyzer")
        for i, p in enumerate(selected):
            _, _, note = estimate_cpe(v.tau, v.gamma, p, v.frequency)
            for col in (7+4*i, 8+4*i):
                ws.cell(n+1, col).comment = Comment(note, "DRT Analyzer")
        title = re.sub(r"[\[\]:*?/\\]", "_", raw.name).strip("'")
        sheet = wb.create_sheet(f"{n:03d} {title}"[:31])
        blocks = [
            (1, ["Experimental f (Hz)", "Experimental Re Z (Ohm)", "Experimental -Im Z (Ohm)", "Imaginary original (Ohm)"],
             zip(raw.frequency, raw.real, -raw.impedance.imag, raw.imaginary_original)),
            (6, ["Fit f (Hz)", "Fit Re Z (Ohm)", "Fit -Im Z (Ohm)"],
             zip(v.frequency, v.reconstructed.real, -v.reconstructed.imag)),
            (10, ["DRT tau (s)", "DRT f (Hz)", "DRT gamma (Ohm per d ln tau)"],
             zip(v.tau, 1/(2*np.pi*v.tau), v.gamma))]
        for start, headers, data in blocks:
            for j, value in enumerate(headers, start):
                sheet.cell(1, j, value)
            for row, values in enumerate(data, 2):
                for col, value in enumerate(values, start):
                    sheet.cell(row, col, float(value))
        sheet.cell(1, 1).comment = Comment(
            f"{raw.name}\nSHA-256: {raw.source_hash}\nVariant: {v.id}\n"
            "Experimental block contains all original points. Fit contains only analysis points. "
            "Blocks have independent frequency/time axes; rows are not paired across blocks. "
            f"Import convention: {raw.info.im_convention}; multiplier={raw.info.im_multiplier}.", "DRT Analyzer")
    for sheet in wb.worksheets:
        _format_sheet(sheet)
        if sheet.title != "Summary":
            sheet.auto_filter.ref = None
    wb.save(path)


def settings_rows(values, prefix=""):
    for key, value in values.items():
        name = f"{prefix}.{key}" if prefix else key
        if isinstance(value, dict):
            yield from settings_rows(value, name)
        elif isinstance(value, (list, tuple)):
            for index, item in enumerate(value):
                yield from settings_rows({str(index+1): item}, name)
        else:
            yield name, value


def export_txt_bundle(directory, pairs):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    pairs = list(pairs)
    write_txt(directory/"summary.txt", summary_rows(pairs))
    for index, (raw, v) in enumerate(pairs, 1):
        stem = re.sub(r'[<>:"/\\|?*]', "_", Path(raw.name).stem)[:70]
        prefix = f"{index:03d}_{stem}_{v.id[:8]}"
        for name, rows in detail_tables(raw, v).items():
            write_txt(directory/f"{prefix}_{name.replace(' ', '_')}.txt", rows)
        write_txt(directory/f"{prefix}_settings.txt", [["Parameter", "Value"], *settings_rows(metadata(raw, v))])
