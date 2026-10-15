"""One file = one spectrum. No row deletion, duplicate merging or silent sign repair."""
import csv
import hashlib
import io
import re
from pathlib import Path
import numpy as np
from .column_detection import ImportErrorDetail, detect_columns, imaginary_sign
from .models import ImportInfo, RawSpectrum


def number(value):
    if isinstance(value, (int, float)):
        return float(value)
    return float(str(value).strip().replace("−", "-").replace("\u00a0", "").replace(",", "."))


def _text_rows(blob):
    if blob.startswith((b'\xff\xfe\x00\x00', b'\x00\x00\xfe\xff')):
        encodings = ("utf-32",)
    elif blob.startswith((b'\xff\xfe', b'\xfe\xff')):
        encodings = ("utf-16",)
    else:
        encodings = ("utf-8-sig", "cp1251", "cp1252")
    for encoding in encodings:
        try:
            text = blob.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportErrorDetail("Unsupported text encoding; save as UTF-8.")
    lines = text.splitlines()
    # Try each actual separator; decimal commas are never treated as delimiters
    # unless the same header also identifies all required columns as CSV fields.
    for rownum, line in enumerate(lines[:100], 1):
        if not line.strip():
            continue
        for delimiter in (";", "\t", ",", "whitespace", "space"):
            if delimiter == "whitespace":
                split = lambda s: re.split(r"\s{2,}|\t+", s.strip())
            elif delimiter == "space":
                # Keep unit parentheses attached to a header token before
                # splitting single-space tables (Frequency (Hz), Re (Ohm)).
                split = lambda s: re.split(r"\s+", re.sub(r"\s+(\([^)]*\)|\[[^]]*\])", r"\1", s.strip()))
            else:
                split = lambda s, d=delimiter: next(csv.reader([s], delimiter=d))
            headers = split(line)
            try:
                indices = detect_columns(headers)
            except ImportErrorDetail:
                continue
            rows = [(n, split(s)) for n, s in enumerate(lines[rownum:], rownum + 1) if s.strip()]
            return headers, indices, rows, encoding, delimiter, rownum
    # Standard DRTtools export: exactly three numeric columns, without a header.
    # Do not infer positions for arbitrary extra columns (e.g. an index column).
    nonempty = [(n, s) for n, s in enumerate(lines, 1) if s.strip()]
    if nonempty:
        for delimiter in (";", "\t", ",", "space"):
            split = (lambda s: re.split(r"\s+", s.strip())) if delimiter == "space" else (
                lambda s, d=delimiter: next(csv.reader([s], delimiter=d)))
            first = split(nonempty[0][1])
            try:
                valid_first = len(first) == 3 and all(np.isfinite(number(v)) for v in first)
            except (ValueError, TypeError):
                valid_first = False
            if valid_first:
                rows = [(n, split(s)) for n, s in nonempty]
                for rownum, row in rows:
                    if len(row) != 3:
                        raise ImportErrorDetail(f"Headerless DRTtools format requires exactly 3 columns; row {rownum} has {len(row)}.")
                return ["Frequency (Hz)", "Re(Z)", "Im(Z)"], (0, 1, 2), rows, encoding, delimiter, 0
    raise ImportErrorDetail("Cannot identify EIS columns. Use frequency/Re/Im headers, or exactly "
                            "three numeric columns without a header in DRTtools order: frequency, Re(Z), Im(Z).")


def load_spectrum(path, im_multiplier=None):
    path = Path(path).resolve()
    blob = path.read_bytes()
    sheet = ""
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
        try:
            candidates = []
            for ws in wb.worksheets:
                rows = list(ws.iter_rows(values_only=True))
                for i, row in enumerate(rows[:100]):
                    try:
                        indices = detect_columns(row)
                        candidates.append((ws.title, row, indices, list(enumerate(rows[i+1:], i+2)), i+1))
                        break
                    except ImportErrorDetail:
                        pass
            if len(candidates) != 1:
                raise ImportErrorDetail("XLSX must contain exactly one sheet with one identifiable spectrum.")
            sheet, headers, indices, rows, header_row = candidates[0]
            encoding, delimiter = "xlsx", "worksheet"
        finally:
            wb.close()
    elif path.suffix.lower() in {".txt", ".csv", ".tsv"}:
        headers, indices, rows, encoding, delimiter, header_row = _text_rows(blob)
    else:
        raise ImportErrorDetail("Supported files: TXT, CSV, TSV and XLSX.")
    values = []
    for rownum, row in rows:
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        try:
            values.append([number(row[j]) for j in indices])
        except (ValueError, TypeError, IndexError) as exc:
            raise ImportErrorDetail(f"Invalid numeric data at row {rownum}: {row!r}") from exc
    if len(values) < 8:
        raise ImportErrorDetail("At least 8 frequency points are required.")
    data = np.asarray(values)
    if not np.isfinite(data).all():
        raise ImportErrorDetail("NaN or infinity in source data; correct the source explicitly.")
    f, re_z, im_z = data.T
    if np.any(f <= 0):
        raise ImportErrorDetail("Frequency must be strictly positive.")
    if np.unique(f).size != f.size:
        raise ImportErrorDetail("Duplicate frequencies found; averaging is not performed automatically.")
    if np.any(np.hypot(re_z, im_z) == 0):
        raise ImportErrorDetail("Zero complex impedance cannot be modulus-weighted.")
    multiplier, evidence = imaginary_sign(headers[indices[2]], im_z, im_multiplier)
    warnings = []
    if not (np.all(np.diff(f) > 0) or np.all(np.diff(f) < 0)):
        warnings.append("Source frequencies are unsorted; analysis sorts a copy in descending order.")
    gaps = np.diff(np.sort(np.log10(f)))
    if max(gaps) > max(0.75, 3 * np.median(gaps)):
        warnings.append("Large frequency gap: DRT resolution is reduced in that interval.")
    if np.ptp(np.log10(f)) < 1:
        warnings.append("Frequency span below one decade; DRT separation is poorly constrained.")
    info = ImportInfo(encoding, delimiter, tuple(str(headers[j]) for j in indices), indices,
                      "Im(Z)" if multiplier == 1 else "-Im(Z)", multiplier, evidence, sheet, header_row)
    return RawSpectrum(str(path), hashlib.sha256(blob).hexdigest(), f, re_z, im_z, info, tuple(warnings))
